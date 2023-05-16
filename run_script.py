import json
from load_json_into_database import load_database
from os import remove
from docx import Document 
from reusable_data import *


#local variables
exercise_list = []
vocabulary_list = []

try:
    # Specify the path of the file
    path = r'import-files/Unit_1_Lesson_6.docx'
    # read file
    document = Document(path) 
    # Read all tables in word 
    tables = document.tables
except Exception as error:
    print(error)

#Get document header
try:
    head_title = tables[0].cell(1, 3).text.strip()
    head_unit = tables[0].cell(0,4).text.strip()
    head_level = tables[0].cell(0,3).text.strip()
    head_sco = tables[0].cell(0, 1).text.strip()
    head_unit = seprate_unit_number_from_string(head_unit)
except Exception as error:
    print(error)

#start extracting documents contents
for j in range(len(tables)):
    for i in range(0, len(tables[j].rows), 1): 
        try:
            title = tables[j].cell(i, 0).text.strip()
            title = ' '.join(title.split())
            if title == 'Welcome page':
                data_list_1 = []
                question_list_1 = []
                str = tables[j+1].cell(i+2, 1).text.strip()
                str = str.split("\n")
                type = str[0]
                content = str[1]
                data_list_1.append({
                        "type":type,
                        "content":content
                    }
                )
                can_do_statments_separator = tables[j+1].cell(i+1, 1).text.strip()
                can_do_statments_separator = can_do_statments_separator.split("\n")
                data_list_1.append({
                        "type": tables[j+1].cell(i+1, 0).text.strip(),
                        "content": can_do_statments_separator
                    }
                )
                
                exercise_list.append({
                        "type":"Welcome",
                        "data":data_list_1,
                        "questions": question_list_1
                    })
#Sec 1: Main Stimulus 1
            if title == "Main Stimulus 1":                
                data_list_2 = main_stimulas(tables, j, i)
#Sec1: Content of Stimulus 1
            if title == 'Content - text or Dialogue' and tables[j-1].cell(i-1, 0).text.strip() == "Main Stimulus 1":
                question_list_2 = []
                data_list_2 = content_text_dialogue(tables, j , i, data_list_2)
                exercise_list.append({
                        "type":"Stim_1",
                        "data":data_list_2,
                        "questions": question_list_2
                    })
#Sec1: MRW
            if title == "Stim 1 - Set 1 8 to 10 questions":
                data_list_3, question_list_3 = MRW(tables, j, i)
                exercise_list.append({
                    "type":"MRW",
                    "data":data_list_3,
                    "questions": question_list_3
                })
#Sec1: SRI
            if title == "Stim 1 - Set 2 8 to 10 questions":
                data_list_5, question_list_5 = SRI(tables, j, i)
                exercise_list.append({
                        "type":"SR1",
                        "data":data_list_5,
                        "questions": question_list_5
                    })
                
#Sec1: DMC          
            if title == "Stim 1 - Set 3 8 to 10 questions":
                data_list_3, question_list_3 = DMC(tables,j,i)
                img_name = tables[j].cell(i+1, 1).text.strip()
                exercise_list.append({
                    "type":"DMC",
                    "data":data_list_3,
                    "questions": question_list_3
                })

#______________________________________________Main Stimulus 2_________________________________
#Sec 2: Main Stimulus 2 
            if title == "Main Stimulus 2":
                data_list_2 = main_stimulas(tables, j, i)   
#Sec2: Content of Stimulus 2
            if title == 'Content - text or Dialogue' and tables[j-1].cell(i-1, 0).text.strip() == "Main Stimulus 2":
                question_list_2 = [] 
                data_list_2 = content_text_dialogue(tables, j , i, data_list_2)
                exercise_list.append({
                        "type":"Stim_1",
                        "data":data_list_2,
                        "questions": question_list_2
                    })
#Sec2: MRW
            if title == "Stim 2 - Set 1 8 to 10 questions":
                data_list_3, question_list_3 = MRW(tables, j, i)
                exercise_list.append({
                    "type":"MRW",
                    "data":data_list_3,
                    "questions": question_list_3
                })

#Sec2: SRI
            if title == "Stim 2 - Set 2 8 to 10 questions":
                data_list_5, question_list_5 = SRI(tables, j, i)
                exercise_list.append({
                        "type":"SR1",
                        "data":data_list_5,
                        "questions": question_list_5
                    })
                
#Sec2: DMC           
            if title == "Stim 2 - Set 3 8 to 10 questions":
                data_list_3, question_list_3 = DMC(tables,j,i)
                exercise_list.append({
                    "type":"DMC",
                    "data":data_list_3,
                    "questions": question_list_3
                })

#_________________________________________Main Stimulus 3_________________________________________

#Sec 3: Main Stimulus 3
            if title == "Main Stim 3 - Chant":
                data_list_2 = main_stimulas(tables, j, i)   
#Sec 3: Content of Stimulus 3
            if title == 'Content' and "Main Stim 3" in  tables[j-1].cell(i-1, 0).text.strip():
                question_list_2 = [] 
                data_list_2 = content_text_dialogue(tables, j , i, data_list_2)
                exercise_list.append({
                        "type":"Stim_1",
                        "data":data_list_2,
                        "questions": question_list_2
                    })
#Sec 3: SRI
            if title == "Stim 3 - Set 1 8 to 10 questions" or title == "Chant - Set 1 8 to 10 questions":
                data_list_5, question_list_5 = SRI(tables, j, i)
                exercise_list.append({
                        "type":"SR1",
                        "data":data_list_5,
                        "questions": question_list_5
                    })
                
#Vocablulary            
            if title == "Vocabulary":
                image_list = []
                data = tables[j+1].cell(i+3, 1).text.split('\n')
                for index in range(len(data)):
                    try:
                        if "aud" in data[index+1]:
                            vocabulary_list.append({
                                "text":data[index].strip(),
                                "audio":data[index+1].strip(),
                                "image":"",
                            })
                    except Exception as e:
                        pass
                
                data = tables[j+1].cell(i+2, 1).text.split('\n')
                for index in range(len(data)):
                    try:
                        if "png" in data[index] and not "png" in data[index-1]:
                            image_list.append(data[index].strip())
                    except Exception as e:
                        pass
                    
                for object in range(len(vocabulary_list)):
                    try:
                        vocabulary_list[object]['image'] = image_list[object]
                    except Exception as e:
                        pass
        except Exception as error: 
            print(error)
            

#insert entire data and generate json file
try:
    main_dict = {
        "title":head_title,
        "unit":head_unit,
        "level":head_level,
        "sco":head_sco,
        "exercises":exercise_list,
        "vocabulary":vocabulary_list
    }
    file_path = head_sco+'.json'
    # file_path = 'custom_file.json'
    with open(file_path, 'w', encoding='utf-8') as f:
        json.dump(main_dict, f, ensure_ascii=False, indent=4)
        if load_database(file_path):
            print("sucessed.")
        else:
            print('failed.')
except Exception as e:
    print('program failed, expected error.')
    print(e)